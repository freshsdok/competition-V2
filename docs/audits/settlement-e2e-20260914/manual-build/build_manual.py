"""Build a user-facing DOCX from verified screen captures without altering source PNGs.
The optional completion manifest is supplied after real provision/final-confirm screenshots exist.
"""
from pathlib import Path
import argparse
import hashlib
import json
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
AUDIT = BASE.parent
SHOTS = AUDIT / 'screenshots'
FONT = 'Hiragino Sans GB'
DISPLAY_RECORDS = []
VERIFIED_SCREENS = {
    'login_tab': {'filename': '13-v2-login-guidance-fixed.png', 'box': [619, 555, 853, 590], 'width': 4.7},
    'login': {'filename': '14-v2-password-hint-fixed.png', 'box': [619, 479, 853, 596], 'width': 4.7},
    'security': {'filename': '15-security-modal-header-fixed.png', 'box': [576, 112, 901, 337], 'width': 3.7},
}

def font(run, size=11, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.insert(0, rfonts)
    for key in ('eastAsia', 'ascii', 'hAnsi', 'cs'):
        rfonts.set(qn('w:' + key), FONT)
    return run

def para(doc, text='', bold=False, after=7, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = Pt(15)
    font(p.add_run(text), 11, bold)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style='Heading ' + str(level))
    p.paragraph_format.space_before = Pt(0 if level == 1 else 11)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = Pt(25 if level == 1 else 18)
    for r in p.runs: font(r, 19 if level == 1 else 13, True)
    return p

def picture(doc, filename, box, width, caption):
    """Use OOXML srcRect display cropping; original screenshot bytes remain unchanged."""
    path = SHOTS / filename
    with Image.open(path) as im: w, h = im.size
    x0, y0, x1, y1 = box
    assert 0 <= x0 < x1 <= w and 0 <= y0 < y1 <= h
    height = width * (y1 - y0) / (x1 - x0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width), height=Inches(height))
    blip_fill = shape._inline.xpath('.//pic:blipFill')[0]
    rect = OxmlElement('a:srcRect')
    for key, fraction in {'l':x0/w, 't':y0/h, 'r':(w-x1)/w, 'b':(h-y1)/h}.items():
        rect.set(key, str(round(fraction * 100000)))
    stretch = blip_fill.find(qn('a:stretch'))
    if stretch is None: blip_fill.append(rect)
    else: blip_fill.insert(list(blip_fill).index(stretch), rect)
    shape._inline.docPr.set('descr', caption)
    cap = doc.add_paragraph(style='Caption')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.line_spacing = Pt(12)
    font(cap.add_run(caption), 9)
    DISPLAY_RECORDS.append({'source':str(path), 'sha256':hashlib.sha256(path.read_bytes()).hexdigest(), 'source_size':[w,h], 'display_crop_pixels':box, 'display_width_inches':width, 'display_height_inches':height, 'caption':caption})

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,width in zip(t.columns,widths):
        col.width=Inches(width)
    for cell,text,width in zip(t.rows[0].cells,headers,widths):
        cell.width=Inches(width); cell.text=text
    for row in rows:
        cells=t.add_row().cells
        for cell,text,width in zip(cells,row,widths):
            cell.width=Inches(width); cell.text=text
    tblpr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for key in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+key); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    tblpr.append(borders)
    for i,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,val in [('top','100'),('left','130'),('bottom','100'),('right','130')]:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),val);m.set(qn('w:type'),'dxa');margins.append(m)
            tcpr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0F6');tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.18
                for r in p.runs:font(r,10,i==0)
        if i==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
    para(doc,'',after=3)
    return t

def setup(review=False):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.62)
    sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    sec.footer_distance=Inches(.3)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        st=doc.styles[name];st.font.name=FONT;st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.line_spacing=Pt(15)
    doc.styles['Normal'].font.size=Pt(11)
    for grid in sec._sectPr.xpath('./w:docGrid'):
        grid.getparent().remove(grid)
    for node in doc.styles._element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    for rf in doc.styles._element.xpath('.//w:rFonts'):
        for key in list(rf.attrib):
            if 'theme' in key.lower(): del rf.attrib[key]
    for color in doc.styles._element.xpath('.//w:color'):
        for key in list(color.attrib):
            if 'theme' in key.lower(): del color.attrib[key]

    doc.core_properties.title='获奖团队银行卡资料填报操作手册'
    doc.core_properties.subject='本人填写 保存 提供和团队最终确认'
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer='审阅稿 收尾步骤待实测  ·  示例截图使用模拟银行资料  ·  ' if review else '示例截图使用模拟银行资料  ·  '
    font(p.add_run(footer),8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    return doc

def build(output, completion=None, review=False):
    DISPLAY_RECORDS.clear()
    if completion:
        required = {'provided', 'tour', 'completed'}
        missing = required - completion.keys()
        if missing:
            raise ValueError('Final manual requires replacement verified screenshots: ' + ', '.join(sorted(missing)))
    doc=setup(review)
    p=doc.add_paragraph('获奖团队银行卡资料填报操作手册',style='Title')
    for r in p.runs:font(r,23,True)
    p.paragraph_format.space_after=Pt(12)
    p.paragraph_format.line_spacing=Pt(30)
    para(doc,'每队由一名队员办理，使用本人的账号和银行卡。请依次完成填写保存、确认提供和团队最终确认。',after=9)
    if review:
        para(doc,'审阅稿 已实测至资料保存',bold=True,after=7)
        para(doc,'本人登录、密码安全验证、资料保存及重新进入已实际验证。选择单位与勾选要求已检查；资料提供、撤销后续访问、完成引导和团队最终确认尚待实测。未实测步骤及其他操作帮助按现有页面与功能逻辑说明。',after=9)
    para(doc,'示例截图使用模拟银行资料，不能用于实际收款。',after=14)
    heading(doc,'1 进入填报并使用本人账号')
    para(doc,'登录原平台，在个人中心打开“我的收款信息v2”。找到自己的获奖团队，点击“申请办理并进入填报”；已经由您办理的，点击“继续填报”。')
    para(doc,'若显示“队员办理中”，请由页面显示的办理人继续操作。一个团队完成后，其他队员无需重复填报，也不能再从本队入口填报。')
    heading(doc,'本人账号登录',2)
    para(doc,'填报窗口如要求登录，请使用与原平台办理人一致的本人账号。尚未认领原账号的，选择“原账号 / 密码”，使用原账号和原密码，按提示登录和认领。')
    c=VERIFIED_SCREENS['login_tab'];picture(doc,c['filename'],c['box'],c['width'],'选择原账号 / 密码入口')
    c=VERIFIED_SCREENS['login'];picture(doc,c['filename'],c['box'],c['width'],'向下滚动填写用户名与密码 原平台账号使用原密码')
    para(doc,'填写后继续向下滚动，点击“登录并进入平台”。如页面要求认领原账号，按提示完成后继续办理。')
    para(doc,'提示账号不一致时，点击“退出当前账号并重新登录”，再使用本人账号进入。已登录正确账号时，可以直接进入填报。')

    heading(doc,'2 填写银行卡并保存').paragraph_format.page_break_before=True
    para(doc,'填写收款人姓名、收款联系手机号、身份号码，以及开户省份、城市、银行、完整支行和银行卡号。银行卡须属于收款人本人。')
    para(doc,'除CNAPS联行号为选填外，其余项目均需填写。核对无误后点击“保存并继续”。')
    picture(doc,'04-empty-form.png',[475,392,1003,568],6.4,'填写银行卡信息后点击保存并继续')
    heading(doc,'出现安全验证时',2)
    para(doc,'按提示验证当前本人账号。可使用已绑定手机的验证码；没有绑定手机号时，可选择已设置的“当前密码”。验证通过后会继续刚才的操作。')
    if review:
        para(doc,'本稿已实际验证当前密码方式，手机验证码方式尚待实测。')
    elif completion:
        para(doc,'本手册截图展示当前密码方式；手机验证码方式本轮未实际验证。')
    c=VERIFIED_SCREENS['security'];picture(doc,c['filename'],c['box'],c['width'],'未绑定手机号时可以使用当前密码完成安全验证')
    para(doc,'点击“取消”会保留当前页面草稿，但不会继续提交。关闭或刷新页面不会替您保存草稿。收款联系手机号仍须填写，它不会修改账号绑定手机号。',after=0)

    heading(doc,'3 选择指定单位并确认提供').paragraph_format.page_break_before=True
    para(doc,'看到“银行卡已保存”后，请继续选择大赛明确指定的接收单位。不确定应选哪一家时，先联系大赛办理方。')
    if review:
        para(doc,'保存或更新后，接收单位仍需主动选择，同意复选框不会自动勾选。')
        picture(doc,'17-updated-profile-saved.png',[482,241,991,307],6.4,'实际更新保存后 接收单位仍为未选择')
    para(doc,'核对单位名称，勾选“我同意将收款资料提供给……，用于大赛结算。”旁的复选框，再点击“确认提供”。如需安全验证，按提示完成即可。',bold=True)
    picture(doc,'08-unit-selected-consent-required.png',[474,180,1003,476],6.4,'选择单位后仍需勾选同意 才能点击确认提供')
    para(doc,'选择单位不会自动提供资料。切换单位后，原来的勾选会清空，需要重新核对并勾选。')
    para(doc,'可展开“谁可以使用这些资料？”查看使用范围。仅指定单位获授权的财务人员可使用当前资料及后续更新。')
    if completion:
        para(doc,'提供成功后，页面会显示“收款资料已提供”。此时仍需完成最后一步。',bold=True)
        c=completion['provided'];picture(doc,c['filename'],c['box'],c.get('width',6.4),'资料已提供给指定单位 仍需完成团队最终确认')
    elif review:
        para(doc,'待实测范围 点击“确认提供”后的成功结果、安全验证衔接及最终完成引导尚未验证。本页截图只展示已检查的选择单位和勾选要求。',bold=True,after=0)
    else:
        # Internal preview stops at the observed state; no invented success screenshot or completed status.
        para(doc,'提供操作之后，请按页面提示继续，最终完成步骤将在取得实际页面截图后编排。',after=0)

    if completion or review:
        heading(doc,'4 点击顶部按钮完成团队确认').paragraph_format.page_break_before=True
        if review:
            para(doc,'本步骤待实测 以下是当前功能的操作说明，尚无本轮成功引导或最终完成截图。',bold=True)
            para(doc,'按当前流程，提供成功并核对通过后，页面应高亮上方“我已完成填报”按钮，并提示“点击这里，才算办完”。请按该引导找到顶部按钮。')
            para(doc,'点击“我已完成填报”后，需等待页面确认结果。预期出现“本队已完成填报”以及办理人和完成时间时，才表示本队流程完成。若没有成功提示，请先查看错误或刷新状态，不要仅凭点击过按钮判断完成。',bold=True)
        else:
            para(doc,'提供成功并核对通过后，页面会高亮上方“我已完成填报”按钮，并提示“点击这里，才算办完”。')
            c=completion['tour'];picture(doc,c['filename'],c['box'],c.get('width',6.4),'跟随页面引导 点击上方我已完成填报')
            para(doc,'请点击“我已完成填报”，等待出现“本队已完成填报”以及办理人和完成时间。到这一步，才完成本队填报流程。',bold=True)
            c=completion['completed'];picture(doc,c['filename'],c['box'],c.get('width',6.4),'显示本队已完成填报后 本队入口关闭')
        para(doc,'点击“稍后确认，收起引导”只关闭提示，仍需完成团队确认。需要重新找到按钮时，点击“查看完成按钮 ↑”。“已确认完成”不表示已发奖或到账，结算进度以大赛通知为准。')
        table(doc,['页面状态','接下来怎么做'],[
            ['银行卡已保存','选择指定单位 勾选同意并确认提供'],
            ['收款资料已提供','点击顶部我已完成填报'],
            ['本队已完成填报','等待大赛后续结算通知']], [2.15,4.75])

    h=heading(doc,'中断办理和常见问题')
    h.paragraph_format.page_break_before=not review
    help_items=[
        ('入口过期或需要重新进入','每次入口有效30分钟。回原平台点击“继续填报”，或在本队“更多操作”中点击“重新进入填报”。重新进入后，需要核对已保存资料并重新选择接收单位。'),
        ('中途关闭页面','关闭页面不会释放办理名额。已经保存的资料可以重新查看；尚未保存的内容需要重新填写。'),
        ('顶部完成按钮不能点击','先确认下方已成功“确认提供”。正在核对时请稍候；有失败提示时，可点击“刷新状态”，或在收款区域“更多操作”中选择“更新已保存状态”，再核对单位。'),
        ('最终确认前需要修改','点击“修改收款信息”，重新填写完整资料并保存，再按提示核对提供状态。已确认完成后需要更正的，请联系大赛办理方。'),
        ('找不到接收单位','点击“重新加载单位”。仍找不到时联系大赛办理方，不要改选与本次大赛无关的单位。'),
        ('身份号码无法核对','此提示表示本次提交未生效。先检查是否填入本人证件上的完整身份号码；填写无误时，按页面提示通过身份处理入口核实，再回来办理。当前页面的草稿会保留。'),
        ('希望其他队员接手','停止填写并关闭其他填报窗口，在本队“更多操作”中点击“放弃本次办理，释放名额”，按确认框操作。放弃不会删除已保存或提供的资料。'),
        ('希望停止后续资料访问','在收款区域“更多操作”的提供记录中选择“撤销后续访问”。撤销不删除已提交的历史结算记录和已合法导出的资料。'),
    ]
    for index,(label,body) in enumerate(help_items):
        if review and index==3:
            heading(doc,'资料修改和异常处理').paragraph_format.page_break_before=True
        heading(doc,label,2);para(doc,body)
        if review and label=='身份号码无法核对':
            picture(doc,'16-identity-conflict-guidance-fixed.png',[474,148,1003,198],6.4,'实际核对失败提示说明 本次提交未生效并引导检查本人证件')
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    records_json=json.dumps(DISPLAY_RECORDS,ensure_ascii=False,indent=2)+'\n'
    (BASE/'image-display-records.json').write_text(records_json)
    (BASE/(output.stem+'-image-display-records.json')).write_text(records_json)
    print(json.dumps({'output':str(output),'images':len(DISPLAY_RECORDS),'completion_included':bool(completion),'review_draft':review},ensure_ascii=False))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=BASE/'layout-preview.docx');parser.add_argument('--completion',type=Path);parser.add_argument('--review',action='store_true')
    args=parser.parse_args();completion=json.loads(args.completion.read_text()) if args.completion else None
    build(args.output.resolve(),completion,args.review)
